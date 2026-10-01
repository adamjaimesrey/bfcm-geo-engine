"""Stage 5 - LLM copy-remediation (v2). One .docx of paste-ready GEO edits per
priority post, via the Anthropic Messages API.

Encodes the Stage-5 rules (see reference/stage5-prompt-issues.md):
  A products = categories the page discusses (add up to 2 per category)
  B paste-ready copy (editorial=placement only)   O tables max 5 columns
  P meta title <=60 / description <=155 chars      Q tables max 5 rows (representative subset)
  C persona woven into table+FAQ (no standalone gifting section)
  F no non-Ecommerce product type as a comparison row
  J AirLift suction-steamer KW-volume caveat (warn)
  K 3-way alignment: meta+H1 <-> body <-> table (fail loud if impossible)
  L native category naming (reference/naming-conventions.md); no "traditional"
  M spec-file routing by category (garment/fans/vacuums); specs verbatim, never invent
  N no comparison table when <2 same-category products (still do the rest)

  R refocus (focus-keywords.csv, refocus=YES only): meta/H1/answer-first + 4 fan-out sections
  S no contradictions (answer-first vs data)       T warnings give exact current -> replacement text
Requires ANTHROPIC_API_KEY in .env.

Usage:
  python scripts/remediate.py            # all rows in the worklist
  python scripts/remediate.py <slug>     # single page (cheap test, e.g. airlift-suction-vs-steam-iron)
"""

import csv
import json
import os
import sys
import time
from pathlib import Path

import pandas as pd
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Pt, RGBColor
from anthropic import Anthropic

BASE_DIR   = Path(__file__).resolve().parent.parent
WORKLIST   = BASE_DIR / "output" / "sihs-worklist.xlsx"
HTML_DIR   = BASE_DIR / "data" / "raw" / "html"
REF        = BASE_DIR / "reference"
SPEC_FILES = {"garment": REF / "product-spec-garment.csv",
              "fans":    REF / "product-spec-fans.csv",
              "vacuums": REF / "product-spec-vacuums.csv"}
NAMING     = REF / "naming-conventions.md"
FOCUS_FILE = REF / "focus-keywords.csv"
OUT_DIR    = BASE_DIR / "output" / "edits"
ENV_FILE   = BASE_DIR / ".env"
MODEL      = "claude-sonnet-5"   # if this 404s, swap to your current model string
MAX_TOKENS = 32000   # room for reasoning + the full answer
RED        = RGBColor(0xC0, 0x39, 0x2B)
AMBER      = RGBColor(0xB0, 0x6A, 0x00)


def load_api_key():
    if ENV_FILE.exists():
        for line in ENV_FILE.read_text().splitlines():
            line = line.strip()
            if line.startswith("ANTHROPIC_API_KEY") and "=" in line:
                return line.split("=", 1)[1].strip().strip('"').strip("'")
    return os.environ.get("ANTHROPIC_API_KEY")


def page_text(slug, limit=6000):
    f = HTML_DIR / f"{slug}.html"
    if not f.exists():
        return ""
    soup = BeautifulSoup(f.read_text(encoding="utf-8", errors="ignore"), "lxml")
    for t in soup(["script", "style", "nav", "footer", "header"]):
        t.decompose()
    return " ".join(soup.get_text(" ").split())[:limit]


def page_meta(slug):
    """Current meta title, meta description and H1 from the cached HTML (the body excerpt strips <head>)."""
    f = HTML_DIR / f"{slug}.html"
    if not f.exists():
        return "=== CURRENT META ===\n(no cached HTML)\n\n"
    soup = BeautifulSoup(f.read_text(encoding="utf-8", errors="ignore"), "lxml")
    title = soup.title.get_text(" ", strip=True) if soup.title else ""
    md = soup.find("meta", attrs={"name": "description"})
    desc = (md.get("content") or "").strip() if md else ""
    h1 = soup.find("h1")
    h1 = h1.get_text(" ", strip=True) if h1 else ""
    return (f"=== CURRENT META ===\nMeta title: {title or '(none)'}\n"
            f"Meta description: {desc or '(none)'}\nH1: {h1 or '(none)'}\n\n")


def read_text(p):
    return p.read_text(encoding="utf-8", errors="ignore") if p.exists() else ""


def spec_data():
    """All three spec files as CSV text, labelled. The prompt routes per Issue M."""
    blocks = []
    for name, f in SPEC_FILES.items():
        if f.exists():
            blocks.append(f"=== SPEC FILE: {name} ({f.name}) ===\n{read_text(f)}")
    return "\n\n".join(blocks)


def load_focus():
    """Refocus targets (Issue I): only rows marked refocus=YES, keyed by slug."""
    if not FOCUS_FILE.exists():
        return {}
    out = {}
    for r in csv.DictReader(open(FOCUS_FILE, encoding="utf-8-sig")):
        if (r.get("refocus") or "").strip().upper() == "YES":
            out[(r.get("slug") or "").strip().lower()] = r
    return out


def focus_block(meta):
    fr = meta.get("focus")
    if not fr:
        return ""
    lines = ["=== REFOCUS TARGET (rule R) ===",
             "Focus keyword: " + (fr.get("focus_keyword") or "").strip(),
             "Focus prompt: " + (fr.get("focus_prompt") or "").strip(),
             "Query fan-outs:"]
    for i in range(1, 5):
        lines.append(f"  {i}. " + (fr.get("fan_out_" + str(i)) or "").strip())
    lines.append("Connected entities: " + (fr.get("connected_entities") or "").strip())
    return "\n".join(lines) + "\n\n"


def build_prompt(meta, naming, specs):
    return f"""You are a senior SEO/GEO editor producing a PASTE-READY edit brief for a Ecommerce
blog post. The client's content editor pastes your output onto the page with ZERO writing of their own.

PAGE: {meta['slug']}
URL: {meta['url']}
Content type: {meta['content_type']}   Category: {meta['category']}   GEO score: {meta['geo_score']}

=== NAMING CONVENTIONS (canonical category names - use verbatim in H1/H2/table headers) ===
{naming}

=== PRODUCT SPEC DATA (specs are verbatim truth; positive terms like "Instant heat" already applied) ===
{specs}

{page_meta(meta['slug'])}=== CURRENT PAGE CONTENT (excerpt - identify on-page products, current meta/H1, body categories) ===
{page_text(meta['slug'])}

{focus_block(meta)}RULES (follow ALL exactly):
- A. PRODUCT SCOPE = CATEGORIES THE PAGE DISCUSSES: find every product category the body discusses under
  any name (canonical, allowed variant, or off-vocabulary term like "traditional steamer"), mapped via the
  NAMING CONVENTIONS. Include every product already on the page. For each discussed category with fewer
  than 2 products on the page, ADD products from the SPEC DATA for that category until it has 2 - choose the
  best fit for this article's angle, ideally at different price/use points. Every added product must appear
  in BOTH the comparison table AND the body: give a paste-ready body sentence naming it (with its spec-file
  URL) and its placement in "added_product_copy", and list it in "warnings" as "Added, not on current page:
  <product>". Exceptions: never add a category the body does not discuss; never add accessories (ironing
  boards, lint removers, cleaning kits) as comparison products; never add non-Ecommerce types (rule F); added
  products must exist in the SPEC DATA, never invented (if a category has fewer than 2, use what exists and
  warn); products that differ only in colour (ALL spec values identical in the SPEC DATA) count as ONE product and share ONE row; if any spec value differs, treat them as separate products (e.g. "AirLift Spin /
  Spin Deluxe"). Pull every product's specs from the SPEC DATA (rule M): route by category to the matching
  spec file, values verbatim, NEVER invent a spec or column, NEVER write "N/A" or "Not specified". If an
  on-page product is missing from the SPEC DATA, add a "warnings" item and use only what the page states.
- B. PASTE-READY: every field except editorial_notes must be FINISHED copy - the exact words for the page.
  editorial_notes = PLACEMENT ONLY (where to paste each block), never what to write.
- C. PERSONA WOVEN IN: no standalone gifting section. Weave recipient/gift persona into (1) the comparison
  table's "Best For (who it suits)" column and (2) one dedicated gift/persona FAQ.
- F. NO NON-ECOMMERCE COMPARATORS: never place a product type Ecommerce does not sell as a table row beside
  Ecommerce products. A competitor category may be named in body prose only, never as a table row.
- J. AIRLIFT CAVEAT: if the page references AirLift / suction steamers, add a "warnings" item:
  "Expect little-to-no KW volume - suction steamers are new to market; do not let KW volume drive
  optimization here."
- K. THREE-WAY ALIGNMENT: the product-category SET must be identical across (1) meta title+description+H1,
  (2) body, (3) comparison-table rows. If the current page is misaligned, propose corrected
  proposed_meta_title / proposed_meta_description / proposed_h1 so all three match, and build the table to
  that same set. If you cannot make all three align, add the specific conflict to "warnings" (fail loud).
- L. NATIVE NAMING: use ONLY the canonical category names above in proposed_h1, any H2s, and table headers.
  Never "traditional iron/steamer" or other off-vocabulary terms (map "traditional iron" -> "steam iron").
  Add any deviation you had to correct to "warnings".
- N. COMPARISON-TABLE THRESHOLD: after applying rule A, if there are fewer than 2 products (rows) to compare
  (e.g. a stick-vacuum page), set comparison_table to null and still produce everything else. Never
  fabricate a comparison.
- O. TABLE WIDTH: max 5 columns - "Product" (include its category, e.g. "Brisa Cordless (Steam Iron)") +
  "Best For (who it suits)" + up to 3 spec columns that apply to every row and best separate the options for
  this article. Where a spec does not apply to a product type, use the spec file's positive term.
- P. META LENGTH: proposed_meta_title max 60 characters (including spaces and any "| Ecommerce"); proposed
  meta description max 155 characters. Count before answering; shorten rather than exceed.
- Q. TABLE ROWS: max 5 rows. If more products are in scope, show the most relevant ones for this article
  while covering every discussed category; products not shown in the table still stay in the body.
- R. REFOCUS (applies ONLY if a REFOCUS TARGET block appears above): refocus the page on that focus keyword.
  proposed_meta_title, proposed_meta_description and proposed_h1 must target the focus keyword (still obeying
  K, L, P). answer_first must directly answer the focus prompt. In "fanout_answers", give one paste-ready new
  body section per query fan-out (all 4): a heading in canonical naming, a 60-120 word answer, and placement.
  Cover every connected entity somewhere in the new copy; list any you could not cover in "warnings". Do not
  repeat fan-out answers in the FAQ block. If no REFOCUS TARGET block appears, set fanout_answers to null.
- S. NO CONTRADICTIONS: answer_first (and any intro/summary copy) must be fully consistent with every spec,
  table row, fan-out section and FAQ in this brief. Never state a benchmark or threshold (e.g. "look for
  30+ g/min") that a Ecommerce product recommended on the page does not meet. Check before answering.
- T. SPOON-FED WARNINGS: any warning that requires a copy change on the page must quote the EXACT current
  text and give the EXACT replacement, e.g. 'Replace "1600 W power" with "1500 W power"'. Never write
  "recommend updating X" without the replacement wording. If several fixes are possible, choose the best ONE
  and give its exact text - never offer options.

Return ONLY a JSON object (no prose, no code fences). Populate a key only where it improves the page; use
null to skip. All content verbatim/paste-ready.

{{
  "warnings": ["fail-loud items: alignment conflicts, naming corrections, AirLift KW caveat, missing specs. Copy fixes must quote exact current text -> exact replacement (rule T). Empty list if none."],
  "proposed_meta_title": "corrected meta title (max 60 characters) if the current one is misaligned/off-vocabulary/too long; else null",
  "proposed_meta_description": "corrected meta description (max 155 characters); else null",
  "proposed_h1": "corrected H1 (canonical naming); else null",
  "answer_first": "40-60 word answer-first paragraph to insert under H1",
  "comparison_table": {{"caption": "one line naming what the table is", "headers": ["..."], "rows": [["..."]]}} or null,
  "added_product_copy": [{{"product": "...", "url": "...", "sentence": "paste-ready body sentence naming the product", "placement": "where in the body it goes"}}],
  "fanout_answers": [{{"fan_out": "...", "heading": "...", "answer": "...", "placement": "..."}}] or null,
  "faq_block": [{{"q": "...", "a": "..."}}],
  "schema_recommendation": "concrete JSON-LD string (Article + BreadcrumbList + FAQPage; add Product/ItemList where products are compared)",
  "editorial_notes": "PLACEMENT ONLY - where each block goes; never what to write"
}}"""


def call_claude(client, meta, naming, specs):
    with client.messages.stream(
        model=MODEL, max_tokens=MAX_TOKENS,
        messages=[{"role": "user", "content": build_prompt(meta, naming, specs)}],
    ) as stream:
        msg = stream.get_final_message()
    raw = "".join(b.text for b in msg.content if getattr(b, "type", "") == "text").strip()
    if not raw:
        raise RuntimeError(f"no answer text (stop_reason={msg.stop_reason}, "
                           f"output_tokens={msg.usage.output_tokens}) - model used its whole budget reasoning")
    if raw.startswith("```"):
        raw = raw.split("```", 2)[1].lstrip("json").strip() if raw.count("```") >= 2 else raw.strip("`")
    return raw


def _h(doc, text, color=RED):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = True
    r.font.size = Pt(13); r.font.color.rgb = color


def render_docx(meta, payload, out_path):
    doc = Document()
    t = doc.add_paragraph(); tr = t.add_run(f"GEO edit brief - {meta['slug']}")
    tr.bold = True; tr.font.size = Pt(18); tr.font.color.rgb = RED
    doc.add_paragraph(f"URL: {meta['url']}")
    if meta.get("semji"):
        doc.add_paragraph(f"SEMJI: {meta['semji']}")
    doc.add_paragraph(f"Type: {meta['content_type']}   |   Category: {meta['category']}   |   "
                      f"GEO score: {meta['geo_score']}   |   Priority #{meta['priority_rank']}")
    doc.add_paragraph("")

    fr = meta.get("focus")
    if fr:
        _h(doc, "Refocus target", color=AMBER)
        doc.add_paragraph("Focus keyword: " + (fr.get("focus_keyword") or ""))
        doc.add_paragraph("Focus prompt: " + (fr.get("focus_prompt") or ""))
        for i in range(1, 5):
            doc.add_paragraph(f"Fan-out {i}: " + (fr.get("fan_out_" + str(i)) or ""))
        doc.add_paragraph("Connected entities: " + (fr.get("connected_entities") or ""))
        doc.add_paragraph("")

    warnings = [w for w in (payload.get("warnings") or []) if str(w).strip()]
    if warnings:
        _h(doc, "\u26a0 Warnings / flags (resolve before publishing)", color=AMBER)
        for w in warnings:
            doc.add_paragraph(f"- {w}")
        doc.add_paragraph("")

    n = 1
    if payload.get("proposed_meta_title") or payload.get("proposed_h1") or payload.get("proposed_meta_description"):
        _h(doc, f"{n}. Proposed metadata + H1 (align meta \u2194 H1 \u2194 body)"); n += 1
        if payload.get("proposed_meta_title"):
            doc.add_paragraph(f"Meta title: {payload['proposed_meta_title']}")
        if payload.get("proposed_meta_description"):
            doc.add_paragraph(f"Meta description: {payload['proposed_meta_description']}")
        if payload.get("proposed_h1"):
            doc.add_paragraph(f"H1: {payload['proposed_h1']}")
        doc.add_paragraph("")

    if payload.get("answer_first"):
        _h(doc, f"{n}. Answer-first block (insert directly under H1)"); n += 1
        doc.add_paragraph(payload["answer_first"]); doc.add_paragraph("")

    ct = payload.get("comparison_table")
    if ct and ct.get("headers") and ct.get("rows"):
        _h(doc, f"{n}. {ct.get('caption', 'Comparison table (render as HTML <table>)')}"); n += 1
        table = doc.add_table(rows=1, cols=len(ct["headers"])); table.style = "Light Grid Accent 1"
        for i, hd in enumerate(ct["headers"]):
            c = table.rows[0].cells[i]; c.text = ""; run = c.paragraphs[0].add_run(str(hd)); run.bold = True
        for row in ct["rows"]:
            cells = table.add_row().cells
            for i, val in enumerate(row[:len(ct["headers"])]):
                cells[i].text = str(val)
        doc.add_paragraph("")
    else:
        _h(doc, f"{n}. Comparison table - OMITTED (fewer than 2 same-category products on page)"); n += 1
        doc.add_paragraph("Per Issue N, no comparison table is generated for a single-product page.")
        doc.add_paragraph("")

    added = payload.get("added_product_copy") or []
    if added:
        _h(doc, f"{n}. Body copy for added products (paste into the body)"); n += 1
        for a in added:
            ap = doc.add_paragraph(); ar = ap.add_run(f"{a.get('product','')} "); ar.bold = True
            ap.add_run(f"({a.get('url','')})")
            doc.add_paragraph(f"Copy: {a.get('sentence','')}")
            doc.add_paragraph(f"Placement: {a.get('placement','')}")
        doc.add_paragraph("")

    fans = payload.get("fanout_answers") or []
    if fans:
        _h(doc, f"{n}. New body sections answering the query fan-outs"); n += 1
        for fa in fans:
            hp = doc.add_paragraph(); hr = hp.add_run(fa.get("heading", "")); hr.bold = True
            doc.add_paragraph(fa.get("answer", ""))
            doc.add_paragraph("Answers fan-out: " + fa.get("fan_out", ""))
            doc.add_paragraph("Placement: " + fa.get("placement", ""))
        doc.add_paragraph("")

    faq = payload.get("faq_block")
    if faq:
        _h(doc, f"{n}. High-intent FAQ block (visible Q&A + FAQPage schema; persona woven into the gift Q)"); n += 1
        for item in faq:
            qp = doc.add_paragraph(); qr = qp.add_run(f"Q: {item.get('q','')}"); qr.bold = True
            doc.add_paragraph(f"A: {item.get('a','')}")
        doc.add_paragraph("")

    if payload.get("schema_recommendation"):
        _h(doc, f"{n}. Structured data (JSON-LD to add)"); n += 1
        cp = doc.add_paragraph(str(payload["schema_recommendation"]))
        cp.runs[0].font.name = "Courier New"; cp.runs[0].font.size = Pt(9); doc.add_paragraph("")

    if payload.get("editorial_notes"):
        _h(doc, f"{n}. Editorial notes (placement only)")
        doc.add_paragraph(payload["editorial_notes"])

    doc.save(out_path)


def validate(payload):
    """Hard checks the model can't be trusted to count: add fail-loud warnings."""
    w = list(payload.get("warnings") or [])
    mt = payload.get("proposed_meta_title") or ""
    md = payload.get("proposed_meta_description") or ""
    if len(mt) > 60:
        w.append(f"Meta title is {len(mt)} characters (max 60) - shorten before publishing.")
    if len(md) > 155:
        w.append(f"Meta description is {len(md)} characters (max 155) - shorten before publishing.")
    ct = payload.get("comparison_table") or {}
    if len(ct.get("headers") or []) > 5:
        w.append(f"Comparison table has {len(ct['headers'])} columns (max 5).")
    if len(ct.get("rows") or []) > 5:
        w.append(f"Comparison table has {len(ct['rows'])} rows (max 5).")
    payload["warnings"] = w
    return payload


def main():
    key = load_api_key()
    if not key:
        sys.exit("FATAL: no ANTHROPIC_API_KEY in .env or environment.")
    if not WORKLIST.exists():
        sys.exit(f"FATAL: {WORKLIST} not found (run Stage 4 first).")

    naming = read_text(NAMING)
    specs = spec_data()
    focus = load_focus()
    if not naming or not specs:
        sys.exit("FATAL: missing reference/naming-conventions.md or spec files.")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    client = Anthropic(api_key=key)
    work = pd.read_excel(WORKLIST)

    only = sys.argv[1].strip().lower() if len(sys.argv) > 1 else None
    ok, failed, consec = [], [], 0
    for _, row in work.iterrows():
        meta = {k: row[k] for k in ["priority_rank", "slug", "url", "category",
                                    "content_type", "geo_score"] if k in row}
        meta.setdefault("url", f"https://www.ecomusa.example/blog/post/{meta['slug']}")
        if "semji" in row:
            meta["semji"] = row["semji"]
        if meta["slug"].lower() in focus:
            meta["focus"] = focus[meta["slug"].lower()]
        if only and meta["slug"].lower() != only:
            continue
        rank = int(meta["priority_rank"]); slug = meta["slug"]
        try:
            raw = call_claude(client, meta, naming, specs)
            payload = validate(json.loads(raw))
            fr = meta.get("focus")
            if fr and (fr.get("focus_intent") or "").strip().lower() != "informational":
                payload["warnings"].append("Refocus: focus keyword intent is not informational - choose an informational alternative.")
            if fr and len(payload.get("fanout_answers") or []) < 4:
                payload["warnings"].append("Refocus: fewer than 4 fan-out answers returned.")
            out = OUT_DIR / f"{rank:02d}_{slug}.docx"
            render_docx(meta, payload, out)
            wcount = len([w for w in (payload.get("warnings") or []) if str(w).strip()])
            print(f"  OK  #{rank:02d} {slug} -> {out.name}  (warnings: {wcount})")
            ok.append(slug); consec = 0
        except json.JSONDecodeError:
            (OUT_DIR / f"{rank:02d}_{slug}.RAW.txt").write_text(raw if "raw" in dir() else "")
            print(f"  PARSE-FAIL #{rank:02d} {slug} (raw saved)"); failed.append(slug); consec += 1
        except Exception as e:
            print(f"  API-FAIL #{rank:02d} {slug}: {e}"); failed.append(slug); consec += 1
        if consec >= 2:
            print("\nTWO-STRIKE STOP: 2 consecutive failures. Halting per CLAUDE.md."); break
        time.sleep(1)

    print(f"\nDone. {len(ok)} docx written, {len(failed)} failed.")
    if failed:
        print("Failed:", failed)


if __name__ == "__main__":
    main()
