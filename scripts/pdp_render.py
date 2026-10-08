"""P5 renderer - turn a brief JSON written by Claude Code into the client .docx (SPEC 9.4/9.7/9.8).
Reads data/processed/pdp-briefs/<group>.json plus that group's score file and captured page.
Checks every 'current' quote appears word-for-word on the page; writes
output/pdp-edits/<priority>_<group>.docx and updates reference/pdp-index.csv.
Usage: python scripts/pdp_render.py --group <group>   |   --all     (exit code 2 if any brief is invalid)"""
import argparse
import csv
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BASE = Path(__file__).resolve().parent.parent
INDEX = BASE / "reference" / "pdp-index.csv"
RAW = BASE / "data" / "raw" / "pdp"
SCORES = BASE / "data" / "processed" / "pdp-scores"
BRIEFS = BASE / "data" / "processed" / "pdp-briefs"
OUT = BASE / "output" / "pdp-edits"
QUESTIONS = [("isA", "What is it?"), ("used_for_audience", "Who is it for / not for?"),
             ("used_for_function", "What job does it do?"), ("used_for_event", "Which situation?"),
             ("capable_of", "What can it do, and limits?"), ("cause_benefit", "Benefit, and why?")]
LABEL = dict(QUESTIONS)
RED, AMBER, GREY, GREEN = RGBColor(0xC0, 0x39, 0x2B), RGBColor(0xB0, 0x6A, 0x00), RGBColor(0x5F, 0x63, 0x68), RGBColor(0x1E, 0x7B, 0x34)
ACTIONS = {"replace", "add", "remove"}


def squash(t):
    return re.sub(r"\s+", " ", str(t)).strip()


def page_text(group):
    d = json.load(open(RAW / group / "sections.json", encoding="utf-8"))
    return squash(" ".join([d.get("h1", "")] + [s["heading"] + " " + s["text"] for s in d["sections"]]))


def validate(brief, group):
    errs, hay = [], page_text(group)
    if brief.get("variant_group") != group:
        errs.append("variant_group does not match file name")
    edits = brief.get("edits", [])
    if not edits:
        errs.append("no edits")
    for i, e in enumerate(edits, 1):
        a = e.get("action")
        if a not in ACTIONS:
            errs.append(f"edit {i}: action must be one of {sorted(ACTIONS)}")
        if not squash(e.get("where", "")):
            errs.append(f"edit {i}: 'where' is empty")
        if a in ("replace", "remove"):
            if not squash(e.get("current", "")):
                errs.append(f"edit {i}: {a} needs the exact current text")
            elif squash(e["current"]) not in hay:
                errs.append(f"edit {i}: current text not found word-for-word on the page: \"{squash(e['current'])[:80]}\"")
        if a in ("replace", "add") and not squash(e.get("new", "")):
            errs.append(f"edit {i}: {a} needs paste-ready new text")
        if not squash(e.get("why", "")):
            errs.append(f"edit {i}: 'why' is empty")
        if not isinstance(e.get("questions"), list) or any(q not in dict(QUESTIONS) for q in e.get("questions", [])):
            errs.append(f"edit {i}: 'questions' must be a list of rubric question keys (empty for a pure naming fix)")
    if "warnings" in brief:
        errs.append("'warnings' is no longer used: put the content editor's questions in 'for_client_editor' and data gaps in 'for_owner'")
    for i, a in enumerate(brief.get("for_client_editor", []), 1):
        for k in ("ask", "if_confirmed", "if_not"):
            if not squash(a.get(k, "")):
                errs.append(f"for_client_editor {i}: '{k}' is empty")
    if not isinstance(brief.get("for_owner", []), list):
        errs.append("'for_owner' must be a list of strings")
    gb = brief.get("gifting_bullet", {})
    if not squash(gb.get("text", "")) or not squash(gb.get("where", "")):
        errs.append("gifting_bullet needs 'where' and 'text'")
    if sorted(brief.get("target_scores", {})) != sorted(q for q, _ in QUESTIONS):
        errs.append("target_scores must cover all six questions")
    return errs


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hex_fill)
    tcPr.append(sh)


def para(doc, text, size=10, bold=False, italic=False, color=None, space_after=4):
    p = doc.add_paragraph(); r = p.add_run(text)
    r.font.size = Pt(size); r.bold = bold; r.italic = italic
    if color: r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space_after)
    return p


def label_line(doc, label, text, color=None, italic=False):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
    a = p.add_run(label + " "); a.bold = True; a.font.size = Pt(9.5); a.font.color.rgb = GREY
    b = p.add_run(text); b.font.size = Pt(10); b.italic = italic
    if color: b.font.color.rgb = color


def heading(doc, text, color=RED):
    p = para(doc, text, size=13, bold=True, color=color, space_after=6)
    p.paragraph_format.space_before = Pt(12)


def render(group, brief, score, idx_rows, out_path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.75); sec.top_margin = sec.bottom_margin = Inches(0.7)
    doc.styles["Normal"].font.name = "Arial"; doc.styles["Normal"].font.size = Pt(10)
    r0 = idx_rows[0]
    now = int(r0["cosmo_score"]) if r0["cosmo_score"] else 0
    target = round(sum(brief["target_scores"].values()) / 18 * 100)

    para(doc, f"[PDP] edit brief - {brief.get('product_name', r0['product_name'])}", size=17, bold=True, color=RED, space_after=6)
    if len(idx_rows) == 1:
        label_line(doc, "Page to edit:", idx_rows[0]["url"])
    else:
        para(doc, "Pages to edit (apply the same edits to each):", size=10, bold=True, color=GREY, space_after=2)
        for r in idx_rows:
            para(doc, f"- {r['product_name']}: {r['url']}", size=10, space_after=1)
    label_line(doc, "", f"Type: pdp   |   Category: {r0['category']}   |   COSMO score: {now}/100 -> target {target}/100   |   Priority #{r0['priority']}")

    asks = brief.get("for_client_editor", [])
    if asks:
        heading(doc, "For the client's content editor: check with Ecommerce (optional)", AMBER)
        para(doc, "The edits below are ready to publish as written. These answers would only let us restore or "
                  "strengthen a claim.", size=9.5, italic=True, color=GREY)
        for i, a in enumerate(asks, 1):
            para(doc, f"{i}. {squash(a['ask'])}", size=10, bold=True, space_after=2)
            label_line(doc, "If confirmed:", squash(a["if_confirmed"]))
            label_line(doc, "If not:", squash(a["if_not"]))

    heading(doc, "COSMO scorecard")
    tbl = doc.add_table(rows=1, cols=5); tbl.style = "Table Grid"; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Inches(1.35), Inches(0.5), Inches(0.6), Inches(2.3), Inches(2.25)]
    for i, h in enumerate(["Question", "Now", "Target", "Current text on the page", "Gap"]):
        c = tbl.rows[0].cells[i]; c.text = ""; run = c.paragraphs[0].add_run(h); run.bold = True; run.font.size = Pt(9)
        shade(c, "F2DCDB")
    for q, label in QUESTIONS:
        v = score["questions"][q]
        quote = v["evidence"][0]["quote"] if v["evidence"] else "(not answered on the page)"
        cells = tbl.add_row().cells
        vals = [label, f"{v['score']}/3", f"{brief['target_scores'][q]}/3", squash(quote)[:220], v.get("gap", "") or "-"]
        for i, val in enumerate(vals):
            cells[i].text = ""; run = cells[i].paragraphs[0].add_run(val); run.font.size = Pt(8.5)
            if i == 2 and brief["target_scores"][q] > v["score"]:
                run.bold = True; run.font.color.rgb = GREEN
    tbl.autofit = False
    grid = tbl._tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths[i].inches * 1440)))
    for row in tbl.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = w

    heading(doc, "Layout")
    para(doc, f"Template {r0['layout_template']}. " + brief.get("layout_note", ""), size=10)

    heading(doc, f"Edits ({len(brief['edits'])}, top of the page to the bottom)")
    verb = {"replace": "Replace", "add": "Add", "remove": "Remove"}
    for i, e in enumerate(brief["edits"], 1):
        p = para(doc, f"Edit {i}  ·  {e['where']}  ·  {verb[e['action']]}", size=10.5, bold=True, space_after=2)
        p.paragraph_format.space_before = Pt(8)
        if e["action"] in ("replace", "remove"):
            label_line(doc, "Current:", squash(e["current"]), color=GREY, italic=True)
        if e["action"] in ("replace", "add"):
            label_line(doc, "Replace with:" if e["action"] == "replace" else "Add:", squash(e["new"]), color=GREEN)
        label_line(doc, "Why:", squash(e["why"]))
        qs = [q for q in e.get("questions", []) if q in LABEL]
        if qs:
            label_line(doc, "Improves:", "; ".join(
                f"{LABEL[q]} ({score['questions'][q]['score']} -> {brief['target_scores'][q]})" for q in qs))

    heading(doc, "BFCM gifting bullet")
    label_line(doc, "Where:", brief["gifting_bullet"]["where"])
    label_line(doc, "Add:", brief["gifting_bullet"]["text"], color=GREEN)

    heading(doc, "Before / after images")
    label_line(doc, "Before:", r0.get("current_screenshot") or "-")
    label_line(doc, "After:", r0.get("after_image") or "added at P6")
    doc.save(out_path)
    return now, target


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--group"); ap.add_argument("--all", action="store_true")
    a = ap.parse_args()
    if not (a.group or a.all):
        sys.exit("Use --group <variant_group> or --all")
    rows = list(csv.DictReader(open(INDEX, encoding="utf-8"))); cols = list(rows[0].keys())
    by_group = {}
    for r in rows:
        by_group.setdefault(r["variant_group"], []).append(r)
    todo = [a.group] if a.group else sorted(f.stem for f in BRIEFS.glob("*.json"))
    OUT.mkdir(parents=True, exist_ok=True)
    bad = 0
    for g in todo:
        bf, sf = BRIEFS / f"{g}.json", SCORES / f"{g}.json"
        if g not in by_group or not bf.exists() or not sf.exists():
            print(f"INVALID {g}: missing index row, brief or score file"); bad += 1; continue
        brief, score = json.load(open(bf, encoding="utf-8")), json.load(open(sf, encoding="utf-8"))
        errs = validate(brief, g)
        if errs:
            bad += 1; print(f"INVALID {g}:"); [print(f"   - {e}") for e in errs]; continue
        members = by_group[g]
        out = OUT / f"{int(members[0]['priority'] or 0):02d}_{g}.docx"
        now, target = render(g, brief, score, members, out)
        todo_owner = [squash(x) for x in brief.get("for_owner", []) if squash(x)]
        for m in members:
            m["edit_docx"] = str(out.relative_to(BASE)); m["status"] = "brief-draft" if todo_owner else "brief"
        print(f"  OK  {g}: {len(brief['edits'])} edits, {now} -> {target}/100 -> {out.name}"
              + (f"  [DRAFT: {len(todo_owner)} item(s) for the owner]" if todo_owner else ""))
        for t in todo_owner:
            print(f"      FOR OWNER: {t}")
    with open(INDEX, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=cols); w.writeheader(); w.writerows(rows)
    if bad:
        sys.exit(2)


if __name__ == "__main__":
    main()
